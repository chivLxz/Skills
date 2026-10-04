#!/usr/bin/env python3
"""Build a Word document with inline screenshots from a JSON spec.

Usage:
    python build_docx.py data.json -o out.docx

Schema (see references/word-output-spec.md):
    {
      "title": "...",
      "subtitle": "...",
      "meta": [{"label": "调研范围", "value": "..."}],
      "sections": [
        {
          "heading": "一、产品形态",
          "intro": "one-sentence finding that opens the section",
          "blocks": [
            {"type": "image", "src": "images/a.png", "caption": "...", "source": "https://..."},
            {"type": "para", "text": "..."},
            {"type": "bullets", "items": ["..."]},
            {"type": "kv", "items": [["标签", "值"]]},
            {"type": "table", "columns": [...], "rows": [[...]]}
          ]
        }
      ],
      "sources": [{"label": "...", "url": "https://...", "accessed": "2026-10-03"}]
    }

Image paths are resolved relative to the JSON file's directory.
"""

import argparse
import json
import os
import sys

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    print("python-docx is required: pip install python-docx", file=sys.stderr)
    sys.exit(1)

BODY_FONT = "Arial"
CJK_FONT = "Microsoft YaHei"
MAX_IMAGE_WIDTH = Inches(6.3)
GRAY = RGBColor(0x59, 0x59, 0x59)
LINK_BLUE = RGBColor(0x05, 0x63, 0xC1)


def set_cjk(run, font=CJK_FONT):
    """python-docx does not set the East Asian font, so do it explicitly."""
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def shade(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_hyperlink(paragraph, url, text=None, size=8.5):
    text = text or url
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(color)
    rpr.append(underline)
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def build(data, out_path, base_dir):
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.25

    _title_block(doc, data)
    _sections(doc, data.get("sections", []), base_dir)
    _sources(doc, data.get("sources", []))

    doc.save(out_path)
    return out_path


def _title_block(doc, data):
    if data.get("title"):
        p = doc.add_heading(level=0)
        run = p.add_run(data["title"])
        set_cjk(run)
    if data.get("subtitle"):
        p = doc.add_paragraph()
        run = p.add_run(data["subtitle"])
        run.font.size = Pt(11)
        run.font.color.rgb = GRAY
        set_cjk(run)

    meta = data.get("meta") or []
    if meta:
        p = doc.add_paragraph()
        run = p.add_run("   |   ".join(
            f"{m.get('label','')}: {m.get('value','')}" for m in meta if m
        ))
        run.font.size = Pt(8.5)
        run.font.color.rgb = GRAY
        set_cjk(run)
    doc.add_paragraph()


def _sections(doc, sections, base_dir):
    for sec in sections:
        if sec.get("heading"):
            h = doc.add_heading(level=1)
            run = h.add_run(sec["heading"])
            set_cjk(run)
        if sec.get("intro"):
            p = doc.add_paragraph()
            run = p.add_run(sec["intro"])
            run.bold = True
            set_cjk(run)
        for block in sec.get("blocks", []):
            _block(doc, block, base_dir)


def _block(doc, block, base_dir):
    btype = block.get("type")

    if btype == "para":
        p = doc.add_paragraph()
        run = p.add_run(block.get("text", ""))
        set_cjk(run)

    elif btype == "bullets":
        for item in block.get("items", []):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(str(item))
            run.font.size = Pt(10.5)
            set_cjk(run)

    elif btype == "kv":
        items = block.get("items", [])
        if not items:
            return
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for label, value in items:
            cells = table.add_row().cells
            cells[0].width = Inches(1.55)
            cells[1].width = Inches(4.7)
            _cell_text(cells[0], str(label), bold=True, bg="F2F2F2")
            _cell_text(cells[1], str(value))
        doc.add_paragraph()

    elif btype == "table":
        cols = block.get("columns", [])
        rows = block.get("rows", [])
        if not cols:
            return
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"
        for i, name in enumerate(cols):
            _cell_text(table.rows[0].cells[i], str(name), bold=True, bg="1F3864",
                       color=RGBColor(0xFF, 0xFF, 0xFF), size=9.5)
        for row in rows:
            cells = table.add_row().cells
            for i in range(len(cols)):
                val = row[i] if i < len(row) else ""
                _cell_text(cells[i], str(val), size=9.5)
        doc.add_paragraph()

    elif btype == "image":
        _image(doc, block, base_dir)


def _cell_text(cell, text, bold=False, bg=None, color=None, size=10):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    set_cjk(run)
    if bg:
        shade(cell, bg)


def _image(doc, block, base_dir):
    src = block.get("src", "")
    path = src if os.path.isabs(src) else os.path.join(base_dir, src)
    if not os.path.isfile(path):
        p = doc.add_paragraph()
        run = p.add_run(f"[截图缺失: {src}]")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        set_cjk(run)
        return

    try:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(path, width=MAX_IMAGE_WIDTH)
    except Exception as exc:
        p = doc.add_paragraph()
        run = p.add_run(f"[截图插入失败: {src} — {exc}]")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        set_cjk(run)
        return

    if block.get("caption"):
        cp = doc.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        crun = cp.add_run(block["caption"])
        crun.font.size = Pt(9)
        crun.font.color.rgb = GRAY
        set_cjk(crun)

    if block.get("source"):
        sp = doc.add_paragraph()
        sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        srun = sp.add_run("来源: ")
        srun.font.size = Pt(8.5)
        srun.font.color.rgb = GRAY
        set_cjk(srun)
        add_hyperlink(sp, block["source"], block["source"], size=8.5)


def _sources(doc, sources):
    if not sources:
        return
    h = doc.add_heading(level=1)
    run = h.add_run("来源清单")
    set_cjk(run)
    for item in sources:
        p = doc.add_paragraph()
        label = item.get("label", "")
        accessed = item.get("accessed", "")
        lrun = p.add_run(f"{label}  " if label else "")
        lrun.bold = True
        lrun.font.size = Pt(9)
        set_cjk(lrun)
        add_hyperlink(p, item.get("url", ""), item.get("url", ""), size=9)
        if accessed:
            arun = p.add_run(f"  访问于 {accessed}")
            arun.font.size = Pt(8.5)
            arun.font.color.rgb = GRAY
            set_cjk(arun)


def main():
    ap = argparse.ArgumentParser(description="Build Word report with screenshots from JSON spec")
    ap.add_argument("data", help="path to the JSON spec")
    ap.add_argument("-o", "--out", default="output.docx")
    args = ap.parse_args()

    try:
        with open(args.data, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, json.JSONDecodeError) as exc:
        print(f"cannot read data file: {exc}", file=sys.stderr)
        return 1

    base_dir = os.path.dirname(os.path.abspath(args.data))
    build(data, args.out, base_dir)
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
