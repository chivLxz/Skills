#!/usr/bin/env python3
"""Build a styled Excel workbook from a JSON spec.

Usage:
    python build_xlsx.py data.json -o out.xlsx

Schema (see references/excel-output-spec.md for the full description):
    {
      "sheets": [
        {
          "name": "核心对比",
          "columns": [{"title": "产品", "width": 22, "wrap": false}],
          "rows": [["产品A", "$20/月"]],
          "notes": "optional footnote line under the table",
          "urls": ["https://example.com"]   # optional -> hyperlink column
        }
      ]
    }

Design notes: header row is bold on a filled background, frozen, and carries an
autofilter so the reader can slice it. Column widths come from the spec because
auto-fit produces unreadably wide columns for long Chinese text.
"""

import argparse
import json
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.hyperlink import Hyperlink
except ImportError:
    print("openpyxl is required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10)
LINK_FONT = Font(name="Arial", size=10, color="0563C1", underline="single")
NOTE_FONT = Font(name="Arial", size=9, italic=True, color="595959")
MISSING_COLOR = "C00000"
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ALIGN_TOP_WRAP = Alignment(vertical="top", wrap_text=True)
ALIGN_TOP = Alignment(vertical="top")
CENTER = Alignment(horizontal="center", vertical="center")


def build(data, out_path):
    wb = Workbook()
    wb.remove(wb.active)

    for sheet_spec in data.get("sheets", []):
        name = (sheet_spec.get("name") or "Sheet")[:31]
        ws = wb.create_sheet(title=name)
        _write_sheet(ws, sheet_spec)

    if not wb.sheetnames:
        ws = wb.create_sheet(title="Empty")
        ws["A1"] = "No sheets defined in the data file."

    wb.save(out_path)
    return out_path


def _write_sheet(ws, spec):
    columns = spec.get("columns") or []
    rows = spec.get("rows") or []
    urls = spec.get("urls") or []

    has_url_col = bool(urls)
    total_cols = len(columns) + (1 if has_url_col else 0)
    if total_cols == 0:
        ws["A1"] = "(empty sheet)"
        return

    # Header
    for idx, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=idx, value=col.get("title", ""))
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

    if has_url_col:
        cell = ws.cell(row=1, column=len(columns) + 1, value="URL")
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    # Body
    for r_off, row in enumerate(rows, start=2):
        for c_off in range(total_cols):
            cell = ws.cell(row=r_off, column=c_off + 1)
            col_spec = columns[c_off] if c_off < len(columns) else {}
            wrap = bool(col_spec.get("wrap"))

            if c_off < len(columns):
                value = row[c_off] if c_off < len(row) else None
            else:
                u_idx = r_off - 2
                value = urls[u_idx] if u_idx < len(urls) else None

            cell.value = value
            cell.font = BODY_FONT
            cell.border = BORDER

            if c_off == len(columns) and has_url_col:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
                if isinstance(value, str) and value.startswith("http"):
                    cell.hyperlink = Hyperlink(ref=cell.coordinate, target=value)
                    cell.font = LINK_FONT
                continue

            cell.alignment = ALIGN_TOP_WRAP if wrap else ALIGN_TOP
            if isinstance(value, str) and value in ("未查到", "未公开", "需联系厂商"):
                cell.font = Font(name="Arial", size=10, color=MISSING_COLOR, italic=True)

    # Column widths
    for idx, col in enumerate(columns, start=1):
        width = col.get("width")
        if width is None:
            longest = max(
                [len(str(col.get("title", "")))]
                + [len(str(r[idx - 1])) for r in rows if idx - 1 < len(r)]
                or [10]
            )
            width = min(max(8, longest + 2), 60)
        ws.column_dimensions[get_column_letter(idx)].width = width

    if has_url_col:
        ws.column_dimensions[get_column_letter(len(columns) + 1)].width = 46

    last_row = max(len(rows) + 1, 1)
    ws.freeze_panes = ws.cell(row=2, column=1)
    if rows:
        ws.auto_filter.ref = (
            f"A1:{get_column_letter(total_cols)}{last_row}"
        )
    ws.row_dimensions[1].height = 26

    note = spec.get("notes")
    if note:
        note_row = last_row + 2
        ws.merge_cells(
            start_row=note_row, start_column=1,
            end_row=note_row, end_column=max(total_cols, 1),
        )
        c = ws.cell(row=note_row, column=1, value=note)
        c.font = NOTE_FONT
        c.alignment = ALIGN_TOP_WRAP
        ws.row_dimensions[note_row].height = 30


def main():
    ap = argparse.ArgumentParser(description="Build Excel workbook from JSON spec")
    ap.add_argument("data", help="path to the JSON spec")
    ap.add_argument("-o", "--out", default="output.xlsx")
    args = ap.parse_args()

    try:
        with open(args.data, "r", encoding="utf-8") as fh:
            data = json.load(fh)
    except (OSError, json.JSONDecodeError) as exc:
        print(f"cannot read data file: {exc}", file=sys.stderr)
        return 1

    build(data, args.out)
    print(f"wrote {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
